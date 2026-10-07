import html
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from prompts import FIELDS
from schema import Card


def chat_text(card: Card) -> str:
    lines = ["📋 <b>Данные документа</b>", ""]
    for field in card.filled:
        mark = "" if field.start >= 0 else " ⚠️"
        lines.append(f"<b>{FIELDS[field.key]}</b>: {html.escape(field.value)}{mark}")
    missing = [FIELDS[f.key] for f in card.fields if not f.value]
    if missing:
        lines += ["", "Не нашёл в документе: " + ", ".join(missing).lower()]
    unchecked = len(card.filled) - card.confirmed
    lines += [
        "",
        f"Сверено с текстом: {card.confirmed} из {len(card.filled)}."
        + (f" ⚠️ — {unchecked} без цитаты, проверь вручную." if unchecked else ""),
    ]
    return "\n".join(lines)


def xlsx(card: Card) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Данные документа"
    sheet.append(["Поле", "Значение", "Откуда в документе", "Сверено с текстом", "Проверил"])
    for field in card.fields:
        sheet.append(
            [
                FIELDS[field.key],
                field.value,
                field.quote if field.value else None,
                ("да" if field.start >= 0 else "нет") if field.value else None,
                None,
            ]
        )
    fill = PatternFill("solid", fgColor="EDEDED")
    for cell, width in zip(sheet[1], (24, 40, 60, 18, 16), strict=True):
        cell.font = Font(bold=True)
        cell.fill = fill
        sheet.column_dimensions[cell.column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
