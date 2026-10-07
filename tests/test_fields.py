import io
import json
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook

import analyzer
import fields
from llm import Reply
from prompts import FIELDS

TEXT = (
    "ДОГОВОР ПОСТАВКИ № 17/24\n"
    "г. Москва, 03.09.2026\n"
    "ООО «Ромашка» (ИНН 7701234567), именуемое Поставщик, и ИП Петров И.И. (ИНН 500100732259), "
    "именуемый Покупатель, заключили договор.\n"
    "2.1. Цена товара составляет 150 000 (сто пятьдесят тысяч) рублей, НДС не облагается.\n"
    "2.2. Оплата производится в течение 5 банковских дней с даты поставки.\n"
)

ANSWER = json.dumps(
    {
        "fields": [
            {"key": "number", "value": "17/24", "quote": "ДОГОВОР ПОСТАВКИ № 17/24"},
            {"key": "date", "value": "03.09.2026", "quote": "г. Москва, 03.09.2026"},
            {"key": "party_a", "value": "ООО «Ромашка»", "quote": "ООО «Ромашка» (ИНН 7701234567)"},
            {"key": "amount", "value": "150 000 руб.", "quote": "Цена товара составляет 150 000"},
            {
                "key": "payment",
                "value": "в течение 5 банковских дней",
                "quote": "придумано моделью",
            },
            {"key": "penalty", "value": None, "quote": None},
            {"key": "лишний", "value": "мусор", "quote": "мусор"},
        ]
    },
    ensure_ascii=False,
)


@pytest.fixture
def model(monkeypatch):
    calls = []

    async def fake_chat(system, content, models, max_tokens=2500):
        calls.append(content)
        return Reply(text=ANSWER, model="test", input_tokens=10, output_tokens=10)

    monkeypatch.setattr(analyzer, "chat", fake_chat)
    return calls


async def test_values_are_anchored_to_their_quotes(model):
    card, _ = await analyzer.extract(TEXT)
    assert [f.key for f in card.fields] == list(FIELDS)
    by_key = {f.key: f for f in card.fields}
    number = by_key["number"]
    assert TEXT[number.start : number.end] == "ДОГОВОР ПОСТАВКИ № 17/24"
    assert by_key["payment"].value and by_key["payment"].start == -1
    assert by_key["penalty"].value is None
    assert len(card.filled) == 5 and card.confirmed == 4
    assert "<<DATA:" in model[0]


async def test_chat_text_flags_unchecked_values(model):
    card, _ = await analyzer.extract(TEXT)
    text = fields.chat_text(card)
    assert "<b>Номер</b>: 17/24" in text
    assert "ООО «Ромашка»" in text
    assert "в течение 5 банковских дней ⚠️" in text
    assert "Сверено с текстом: 4 из 5" in text
    assert "неустойка" in text


async def test_excel_has_source_and_check_columns(model):
    card, _ = await analyzer.extract(TEXT)
    rows = list(load_workbook(io.BytesIO(fields.xlsx(card))).active.iter_rows(values_only=True))
    assert rows[0] == ("Поле", "Значение", "Откуда в документе", "Сверено с текстом", "Проверил")
    amount = next(r for r in rows if r[0] == "Сумма")
    assert amount[1:4] == ("150 000 руб.", "Цена товара составляет 150 000", "да")
    assert next(r for r in rows if r[0] == "Порядок оплаты")[3] == "нет"


def test_escapes_html_from_the_document():
    from schema import Card, DocField

    card = Card(fields=[DocField(key="party_a", value="<b>ООО</b>", quote="x", start=0, end=1)])
    assert "&lt;b&gt;ООО&lt;/b&gt;" in fields.chat_text(card)


async def test_api_returns_fields_as_json(model, monkeypatch):
    from fastapi.testclient import TestClient

    import main

    monkeypatch.setattr(main, "_user", lambda init_data: {"id": 1})
    monkeypatch.setattr(main, "_take", lambda *args: None)
    response = TestClient(main.app).post("/api/fields", json={"init_data": "x", "text": TEXT})
    assert response.status_code == 200
    body = response.json()
    assert body["confirmed"] == 4
    assert body["fields"][0]["label"] == "Тип документа"
    assert {f["key"]: f["value"] for f in body["fields"]}["amount"] == "150 000 руб."


async def test_bot_sends_card_and_excel_in_fields_mode(model, monkeypatch):
    from contextlib import nullcontext

    import bot

    monkeypatch.setattr(
        bot, "ChatActionSender", SimpleNamespace(upload_document=lambda **_: nullcontext())
    )

    message = SimpleNamespace(
        chat=SimpleNamespace(id=5),
        bot=SimpleNamespace(send_chat_action=AsyncMock()),
        answer=AsyncMock(),
        answer_document=AsyncMock(),
    )
    state = SimpleNamespace(get_data=AsyncMock(return_value={"mode": "fields"}))
    await bot._dispatch(message, state, TEXT)
    assert "Данные документа" in message.answer.call_args.args[0]
    document = message.answer_document.call_args.args[0]
    assert document.filename == "Данные документа.xlsx"
